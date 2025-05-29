# utils/file_manager/xlsx.py

import logging
from openpyxl import load_workbook
import re


def extract_text(file_path: str) -> str:
    """
    Excel(.xlsx) 파일의 모든 시트에서 셀 텍스트를 추출하여 하나의 문자열로 반환
    """
    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
        texts = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        cell_str = str(cell).strip()
                        if cell_str:
                            texts.append(cell_str)
        return "\n".join(texts)
    except Exception as e:
        logging.error(f"[XLSX] 텍스트 추출 실패({file_path}): {e}")
        return ""


def extract_first_sentence(file_path: str) -> str:
    """
    Excel(.xlsx) 파일에서 첫 번째 셀의 텍스트를 첫 문장으로 반환
    """
    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        text = str(cell).strip()
                        if not text:
                            continue
                        match = re.match(r".+?[\.\!?](?:\s|$)", text)
                        if match:
                            return match.group(0).strip()
                        return text
        return ""
    except Exception as e:
        logging.error(f"[XLSX] 첫 문장 추출 실패({file_path}): {e}")
        return ""
