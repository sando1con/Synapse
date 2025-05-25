# utils/file_manager/ppt.py
import logging
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import re

def extract_text(file_path: str) -> str:
    """
    PowerPoint(.pptx) 파일의 모든 슬라이드에서 텍스트를 추출하여 하나의 문자열로 반환
    """
    try:
        prs = Presentation(file_path)
        texts = []

        for slide in prs.slides:
            for shape in slide.shapes:
                # 1) 일반 텍스트 프레임
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        if para.text and para.text.strip():
                            texts.append(para.text.strip())

                # 2) 그룹 쉐이프 안의 텍스트도 탐색
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    for shp in shape.shapes:
                        if shp.has_text_frame:
                            for para in shp.text_frame.paragraphs:
                                if para.text and para.text.strip():
                                    texts.append(para.text.strip())

        return "\n".join(texts)

    except Exception as e:
        logging.error(f"[PPT] 텍스트 추출 실패 ({file_path}): {e}")
        return ""

def extract_first_sentence(file_path: str) -> str:
    """
    추출된 텍스트에서 첫 문장(마침표, 물음표, 느낌표로 구분)만 반환
    """
    text = extract_text(file_path)
    if not text:
        return ""

    # 문장 분리: 마침표, 물음표, 느낌표
    sentences = re.split(r'(?<=[\.\?\!])\s+', text)
    return sentences[0].strip() if sentences else ""